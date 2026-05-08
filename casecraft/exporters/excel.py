"""Excel 导出器 - 带样式、筛选、冻结首行"""

from __future__ import annotations

from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from casecraft.core import TestCase

from .base import BaseExporter


PRIORITY_COLORS = {
    "P0": "FF0000",
    "P1": "FF8C00",
    "P2": "4169E1",
    "P3": "808080",
}

TAG_COLORS = {
    "正常流程": "228B22",
    "异常": "FF4500",
    "边界值": "DAA520",
    "权限": "8A2BE2",
}

COL_WIDTHS = {
    "模块": 14, "用例标题": 30, "优先级": 8, "用例类型": 10,
    "前置条件": 25, "操作步骤": 40, "预期结果": 30, "测试数据": 25,
    "关联接口": 25, "标签": 12, "用例来源": 18, "实际结果": 20,
    "状态": 10, "备注": 20, "创建人": 10, "执行人": 10,
}


class ExcelExporter(BaseExporter):
    name = "excel"
    extension = "xlsx"

    def __init__(self, hidden_columns: set[str] | None = None, **kwargs):
        super().__init__(**kwargs)
        self.hidden_columns = hidden_columns or {"前置条件", "测试数据"}

    def export(self, cases: Iterable[TestCase], output_path: str | None = None,
               creator: str = "casecraft", **kwargs) -> str:
        output_path = self._resolve_output(output_path)
        rows = self._normalize(cases, creator)

        wb = Workbook()
        ws = wb.active
        ws.title = "测试用例"

        header_font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_align = Alignment(vertical="top", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )

        for col, header in enumerate(self.headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        for row_idx, case in enumerate(rows, 2):
            for col_idx, header in enumerate(self.headers, 1):
                value = case.get(header, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = cell_align
                cell.border = thin_border

                if header == "优先级" and value in PRIORITY_COLORS:
                    cell.font = Font(color=PRIORITY_COLORS[value], bold=True)
                elif header == "标签" and value in TAG_COLORS:
                    cell.font = Font(color=TAG_COLORS[value])
                elif header == "用例来源" and value:
                    if value.startswith("代码分析"):
                        cell.font = Font(color="0066CC", bold=True)
                        cell.fill = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")
                    elif value == "需求文档":
                        cell.font = Font(color="228B22")

        for col_idx, header in enumerate(self.headers, 1):
            width = COL_WIDTHS.get(header, 15)
            col_letter = _col_letter(col_idx)
            ws.column_dimensions[col_letter].width = width
            if header in self.hidden_columns:
                ws.column_dimensions[col_letter].hidden = True

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{_col_letter(len(self.headers))}1"

        wb.save(output_path)
        return output_path


def _col_letter(idx: int) -> str:
    """1 -> A, 27 -> AA"""
    result = ""
    while idx:
        idx, rem = divmod(idx - 1, 26)
        result = chr(65 + rem) + result
    return result
